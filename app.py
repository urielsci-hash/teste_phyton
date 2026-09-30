from flask import Flask, render_template, request, session, redirect, url_for, jsonify, g, make_response
from flask_session import Session
from werkzeug.security import check_password_hash, generate_password_hash
import os
import secrets
import datetime
import requests
import time
import feedparser
from dotenv import load_dotenv

load_dotenv()

app = Flask(__name__)
app.config['SECRET_KEY'] = os.getenv('SECRET_KEY', 'default_secret_key_change_me')
app.config['SESSION_TYPE'] = 'filesystem'
Session(app)

import database

@app.teardown_appcontext
def close_db_connection(exception):
    # This function runs at the end of each request
    pass

@app.route('/api/clima')
def api_clima():
    import json
    import tempfile

    clima_cidade = os.getenv('CLIMA_CIDADE', 'Serra Negra,BR')
    clima_api_key = os.getenv('CLIMA_API_KEY', '6e5ac0b27e0a520eb089a9f00c2e6556')
    cache_file = os.path.join(tempfile.gettempdir(), 'clima_sinergia_agro.json')

    if os.path.exists(cache_file) and (time.time() - os.path.getmtime(cache_file) < 1800):
        with open(cache_file, 'r', encoding='utf-8') as f:
            return f.read(), 200, {'Content-Type': 'application/json; charset=utf-8'}

    url = f"https://api.openweathermap.org/data/2.5/weather?q={clima_cidade}&units=metric&lang=pt_br&appid={clima_api_key}"
    try:
        r = requests.get(url, timeout=5)
        if r.status_code == 200:
            with open(cache_file, 'w', encoding='utf-8') as f:
                f.write(r.text)
            return r.text, 200, {'Content-Type': 'application/json; charset=utf-8'}
    except:
        pass

    return jsonify({"erro": True})

@app.route('/api/status')
def api_status():
    mes_atual = datetime.datetime.now().strftime('%Y-%m')

    query_status = """
        SELECT s1.data, s1.status
        FROM status_qualidade_dia s1
        INNER JOIN (
            SELECT data, MAX(criado_em) AS max_criado
            FROM status_qualidade_dia
            WHERE data LIKE ?
            GROUP BY data
        ) s2 ON s1.data = s2.data AND s1.criado_em = s2.max_criado
    """
    rows = database.execute_query(query_status, (f"{mes_atual}-%",))
    status_por_dia = {}
    for row in rows:
        dia = int(row['data'].split('-')[2])
        status_por_dia[dia] = row['status']

    query_desvios = """
        SELECT s.data, s.desvio AS descricao_desvio, s.acao_tomada, s.como_evitar, s.observacao
        FROM status_qualidade_dia s
        INNER JOIN (
            SELECT data, MAX(criado_em) AS max_criado
            FROM status_qualidade_dia
            WHERE data LIKE ?
            GROUP BY data
        ) max_s ON s.data = max_s.data AND s.criado_em = max_s.max_criado
        WHERE s.status IN ('atencao', 'grave') AND s.desvio IS NOT NULL AND s.desvio != ''
        ORDER BY s.data DESC, s.id DESC
    """
    desvios = database.execute_query(query_desvios, (f"{mes_atual}-%",))

    return jsonify({
        "status_dias": status_por_dia,
        "desvios": desvios
    })

@app.route('/api/mural')
def api_mural():
    hoje = datetime.datetime.now().strftime('%Y-%m-%d')
    query = """
        SELECT titulo, conteudo, tipo, imagem_path
        FROM mural_posts
        WHERE ativo = 1 AND ? BETWEEN data_inicio AND data_fim
        ORDER BY criado_em DESC
    """
    posts = database.execute_query(query, (hoje,))
    mural = []
    pasta_uploads = os.path.join(os.path.dirname(__file__), 'static', 'uploads', 'mural')

    for post in posts:
        if post['tipo'] == 'imagem':
            existe = post.get('imagem_path') and os.path.exists(os.path.join(pasta_uploads, post['imagem_path']))
            if not existe:
                post['tipo'] = 'texto'
                post['conteudo'] = post.get('conteudo', '')
        mural.append(post)

    return jsonify({"mural": mural})

@app.route('/api/indicadores')
def api_indicadores():
    import json
    hoje = datetime.datetime.now().strftime('%Y-%m-%d')
    query = """
        SELECT nome, tipo_grafico, categorias, valores
        FROM indicadores
        WHERE ativo = 1 AND ? BETWEEN data_inicio AND data_fim
        ORDER BY nome
    """
    rows = database.execute_query(query, (hoje,))
    indicadores = []

    for row in rows:
        try:
            categorias = json.loads(row['categorias']) if row['categorias'] else []
        except:
            categorias = []
        try:
            valores = json.loads(row['valores']) if row['valores'] else []
        except:
            valores = []

        indicadores.append({
            "nome": row['nome'],
            "tipo": row['tipo_grafico'],
            "categorias": categorias,
            "valores": valores
        })

    return jsonify({"indicadores": indicadores})

@app.route('/api/em-formulacao')
def api_em_formulacao():
    hoje = datetime.datetime.now().strftime('%Y-%m-%d')

    query = "SELECT * FROM producao_diaria WHERE data = ?"
    producao = database.execute_query_single(query, (hoje,))

    if not producao:
        query_fallback = "SELECT * FROM producao_diaria ORDER BY data DESC LIMIT 1"
        producao = database.execute_query_single(query_fallback)

    tanques = []
    if producao:
        query_tanques = "SELECT tanque, produto, valor_litros FROM producao_tanques WHERE producao_diaria_id = ?"
        tanques = database.execute_query(query_tanques, (producao['id'],))

    return jsonify({
        "producao": producao,
        "producao_tanques": tanques
    })

@app.route('/api/noticias')
def api_noticias():
    urls = database.execute_query("SELECT url_feed FROM noticias_rss WHERE ativo = 1")
    noticias = []

    for row in urls:
        url = row['url_feed']
        try:
            feed = feedparser.parse(url)
            for entry in feed.entries[:5]: # limit per feed
                noticias.append({
                    "titulo": entry.title,
                    "link": entry.link,
                    "data": entry.get('published', entry.get('updated', '')),
                    "fonte": feed.feed.get('title', url)
                })
        except Exception as e:
            print("Error parsing feed", url, e)

    # Sort by date (string simple sort or using datetime if needed, keeping it simple here)
    # The original PHP limited to 15 globally
    noticias = sorted(noticias, key=lambda x: x['data'], reverse=True)[:15]
    return jsonify({"noticias": noticias})

def is_logged_in():
    return 'usuario_id' in session

def departamento_usuario():
    return session.get('departamento_nome')

def precisa_trocar_senha():
    return session.get('deve_trocar_senha', False)

def get_destino_apos_login(departamento):
    if departamento == 'Visualizacao':
        return url_for('index')
    return url_for('admin_index')

@app.before_request
def auth_middleware():
    if not is_logged_in() and 'lembrar_token' in request.cookies:
        token_hash = request.cookies.get('lembrar_token')

        query = """
            SELECT u.*, d.nome AS departamento_nome
            FROM tokens_lembrar t
            JOIN usuarios u ON u.id = t.usuario_id
            JOIN departamentos d ON d.id = u.departamento_id
            WHERE t.token_hash = ? AND t.expira_em > ? AND u.ativo = 1
        """
        agora = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        usuario = database.execute_query_single(query, (token_hash, agora))

        if usuario:
            session['usuario_id'] = usuario['id']
            session['usuario_nome'] = usuario['nome']
            session['departamento_id'] = usuario['departamento_id']
            session['departamento_nome'] = usuario['departamento_nome']
            session['deve_trocar_senha'] = bool(usuario['deve_trocar_senha'])

    allowed_routes = ['login', 'logout', 'static', 'api_clima', 'api_noticias', 'api_status', 'api_mural', 'api_indicadores', 'api_em_formulacao', 'gerar_senha']
    if request.endpoint not in allowed_routes and not is_logged_in():
        # Exigência de login no index e admin
        return redirect(url_for('login'))

    if is_logged_in() and precisa_trocar_senha() and request.endpoint not in ['trocar_senha', 'logout', 'static']:
        return redirect(url_for('trocar_senha'))

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        email = request.form.get('email', '').strip()
        senha = request.form.get('senha', '')
        lembrar = request.form.get('lembrar')

        query = """
            SELECT u.*, d.nome AS departamento_nome
            FROM usuarios u
            JOIN departamentos d ON d.id = u.departamento_id
            WHERE u.email = ? AND u.ativo = 1
        """
        # Execute parameterized query based on sqlite/mysql adapter
        usuario = database.execute_query_single(query, (email,))

        if usuario and check_password_hash(usuario['senha_hash'], senha):
            session['usuario_id'] = usuario['id']
            session['usuario_nome'] = usuario['nome']
            session['departamento_id'] = usuario['departamento_id']
            session['departamento_nome'] = usuario['departamento_nome']
            session['deve_trocar_senha'] = bool(usuario['deve_trocar_senha'])

            response = make_response(redirect(get_destino_apos_login(usuario['departamento_nome'])))

            if lembrar:
                token = secrets.token_hex(32)
                expira = datetime.datetime.now() + datetime.timedelta(days=365)
                expira_str = expira.strftime('%Y-%m-%d %H:%M:%S')
                database.execute_query("INSERT INTO tokens_lembrar (usuario_id, token_hash, expira_em) VALUES (?, ?, ?)",
                                       (usuario['id'], token, expira_str), fetch=False)
                response.set_cookie('lembrar_token', token, max_age=365*24*60*60)

            return response

        return render_template('login.html', erro="E-mail ou senha inválidos.")

    return render_template('login.html', erro=None)

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

@app.route('/admin')
def admin_index():
    # Simple placeholder for admin redirect based on department
    dep = departamento_usuario()
    if dep == 'Qualidade':
        return redirect(url_for('admin_status'))
    elif dep in ['RH', 'Marketing']:
        return redirect(url_for('admin_mural'))
    elif dep == 'TI':
        return redirect(url_for('admin_usuarios'))
    return redirect(url_for('index'))

@app.route('/admin/status', methods=['GET', 'POST'])
def admin_status():
    if departamento_usuario() != 'Qualidade':
        return "Acesso negado.", 403

    mensagem = None

    anos = list(range(datetime.datetime.now().year - 5, max(2050, datetime.datetime.now().year + 25) + 1))
    meses = {1:"Janeiro",2:"Fevereiro",3:"Março",4:"Abril",5:"Maio",6:"Junho",
             7:"Julho",8:"Agosto",9:"Setembro",10:"Outubro",11:"Novembro",12:"Dezembro"}

    ano_selecionado = int(request.args.get('ano', datetime.datetime.now().year))
    mes_selecionado = int(request.args.get('mes_num', datetime.datetime.now().month))
    dia_pre_selecionado = request.args.get('dia', datetime.datetime.now().strftime('%Y-%m-%d'))

    if request.method == 'POST':
        if 'salvar_status' in request.form:
            data = request.form.get('data')
            status = request.form.get('status')
            observacao = request.form.get('observacao', '').strip() or None

            desvio = None
            acao_tomada = None
            como_evitar = None

            if status in ['atencao', 'grave']:
                desvio = request.form.get('desvio', '').strip() or None
                acao_tomada = request.form.get('acao_tomada', '').strip() or None
                como_evitar = request.form.get('como_evitar', '').strip() or None

            query = """
                INSERT INTO status_qualidade_dia (data, status, observacao, desvio, acao_tomada, como_evitar, usuario_id)
                VALUES (?, ?, ?, ?, ?, ?, ?)
            """
            database.execute_query(query, (data, status, observacao, desvio, acao_tomada, como_evitar, session['usuario_id']), fetch=False)
            registrar_log('qualidade', 'Lançou o status de qualidade de um dia do mês')
            mensagem = "Status do dia salvo."

        elif 'excluir_historico' in request.form:
            reg_id = request.form.get('id')
            database.execute_query("DELETE FROM status_qualidade_dia WHERE id = ?", (reg_id,), fetch=False)
            registrar_log('qualidade', 'Excluiu um registro de status de qualidade')
            mensagem = "Registro excluído."

    mes_str = f"{ano_selecionado}-{mes_selecionado:02d}-%"
    query_historico = """
        SELECT s.*, u.nome AS usuario_nome
        FROM status_qualidade_dia s
        JOIN usuarios u ON u.id = s.usuario_id
        WHERE s.data LIKE ?
        ORDER BY s.data DESC, s.criado_em DESC
    """
    historico_db = database.execute_query(query_historico, (mes_str,))
    historico = []

    # Simple formatting for template
    for r in historico_db:
        dt = r['data']
        criado = str(r['criado_em']) # simplistic representation
        r['data_formatada'] = dt
        r['hora_formatada'] = criado
        historico.append(r)

    query_ultimo = "SELECT * FROM status_qualidade_dia WHERE data = ? ORDER BY criado_em DESC LIMIT 1"
    ultimo_lancamento = database.execute_query_single(query_ultimo, (dia_pre_selecionado,)) or {}

    return render_template('admin_status.html', mensagem=mensagem, anos=anos, meses=meses,
                           ano_selecionado=ano_selecionado, mes_selecionado=mes_selecionado,
                           dia_pre_selecionado=dia_pre_selecionado, historico=historico,
                           ultimo_lancamento=ultimo_lancamento)

@app.route('/admin/indicadores', methods=['GET', 'POST'])
def admin_indicadores():
    if departamento_usuario() != 'Qualidade':
        return "Acesso negado.", 403

    mensagem = None
    erro = None

    if request.method == 'POST':
        if 'enviar_planilha' in request.form:
            import openpyxl
            import json
            arquivo = request.files.get('arquivo')
            data_inicio = request.form.get('data_inicio')
            data_fim = request.form.get('data_fim')

            if arquivo and arquivo.filename.endswith('.xlsx'):
                try:
                    wb = openpyxl.load_wsgi(arquivo) if hasattr(openpyxl, 'load_wsgi') else openpyxl.load_workbook(arquivo, data_only=True)
                    abas = wb.sheetnames
                    if not abas:
                        erro = "Não encontrei nenhuma aba nessa planilha."
                    else:
                        for nome_aba in abas:
                            ws = wb[nome_aba]
                            # row 1 = categorias, row 2 = valores
                            linhas = list(ws.iter_rows(min_row=1, max_row=2, values_only=True))

                            categorias = [str(x) for x in linhas[0] if x is not None] if len(linhas) > 0 else []
                            valores_brutos = linhas[1] if len(linhas) > 1 else []
                            valores = []
                            for i, cat in enumerate(categorias):
                                val = valores_brutos[i] if i < len(valores_brutos) else 0
                                try:
                                    valores.append(float(str(val).replace(',', '.')))
                                except:
                                    valores.append(0.0)

                            if categorias and valores:
                                check = database.execute_query_single("SELECT id FROM indicadores WHERE nome = ?", (nome_aba,))
                                cat_json = json.dumps(categorias, ensure_ascii=False)
                                val_json = json.dumps(valores, ensure_ascii=False)

                                if check:
                                    database.execute_query("UPDATE indicadores SET categorias = ?, valores = ?, data_inicio = ?, data_fim = ?, ativo = 1, usuario_id = ? WHERE id = ?",
                                                           (cat_json, val_json, data_inicio, data_fim, session['usuario_id'], check['id']), fetch=False)
                                else:
                                    database.execute_query("INSERT INTO indicadores (nome, categorias, valores, data_inicio, data_fim, usuario_id) VALUES (?, ?, ?, ?, ?, ?)",
                                                           (nome_aba, cat_json, val_json, data_inicio, data_fim, session['usuario_id']), fetch=False)

                        database.execute_query("INSERT INTO indicadores_arquivos (nome_arquivo, usuario_id) VALUES (?, ?)",
                                               (arquivo.filename, session['usuario_id']), fetch=False)
                        registrar_log('indicadores', 'Fez upload de uma nova planilha de indicadores')
                        mensagem = "Planilha processada com sucesso."
                except Exception as e:
                    erro = "Erro ao ler a planilha: " + str(e)
            else:
                erro = "Envie um arquivo válido no formato .xlsx"

        elif 'alterar_tipo' in request.form:
            ind_id = request.form.get('id')
            tipo = request.form.get('tipo_grafico')
            database.execute_query("UPDATE indicadores SET tipo_grafico = ? WHERE id = ?", (tipo, ind_id), fetch=False)
            registrar_log('indicadores', 'Alterou o tipo de gráfico de um indicador')
            mensagem = "Tipo alterado."

        elif 'alternar_ativo' in request.form:
            ind_id = request.form.get('id')
            database.execute_query("UPDATE indicadores SET ativo = 1 - ativo WHERE id = ?", (ind_id,), fetch=False)
            registrar_log('indicadores', 'Ativou/desativou um indicador')
            mensagem = "Status alterado."

        elif 'excluir' in request.form:
            ind_id = request.form.get('id')
            database.execute_query("DELETE FROM indicadores WHERE id = ?", (ind_id,), fetch=False)
            registrar_log('indicadores', 'Excluiu um indicador')
            mensagem = "Indicador excluído."

    mostrar_todos = request.args.get('mostrar') == 'todos'
    hoje = datetime.datetime.now().strftime('%Y-%m-%d')
    if mostrar_todos:
        indicadores = database.execute_query("SELECT * FROM indicadores ORDER BY nome")
    else:
        indicadores = database.execute_query("SELECT * FROM indicadores WHERE ativo = 1 AND ? BETWEEN data_inicio AND data_fim ORDER BY nome", (hoje,))

    return render_template('admin_indicadores.html', indicadores=indicadores, mensagem=mensagem, erro=erro, hoje=hoje, mostrar_todos='1' if mostrar_todos else '0')

@app.route('/admin/em-formulacao', methods=['GET', 'POST'])
def admin_em_formulacao():
    if departamento_usuario() != 'Qualidade':
        return "Acesso negado.", 403

    mensagem = None
    tanques_disponiveis = ["T1", "T2", "T3", "T4"]
    dia_selecionado = request.args.get('dia', datetime.datetime.now().strftime('%Y-%m-%d'))

    if request.method == 'POST':
        if 'salvar' in request.form:
            data = request.form.get('data')
            try:
                capacidade_total = float(request.form.get('capacidade_total', '0').replace(',', '.'))
            except:
                capacidade_total = 0.0

            armazenamento_utilizado = 0.0
            tanques_data = {}
            for t in tanques_disponiveis:
                val_str = request.form.get(f'tanque_{t}', '').replace(',', '.')
                prod_str = request.form.get(f'tanque_produto_{t}', '').strip()
                if val_str:
                    try:
                        val_float = float(val_str)
                        armazenamento_utilizado += val_float
                        tanques_data[t] = {'valor_litros': val_float, 'produto': prod_str or None}
                    except:
                        pass

            # Insert or update
            check = database.execute_query_single("SELECT id FROM producao_diaria WHERE data = ?", (data,))
            if check:
                prod_id = check['id']
                database.execute_query("UPDATE producao_diaria SET capacidade_total_litros = ?, armazenamento_utilizado_litros = ?, usuario_id = ? WHERE id = ?",
                                       (capacidade_total, armazenamento_utilizado, session['usuario_id'], prod_id), fetch=False)
            else:
                prod_id = database.execute_query("INSERT INTO producao_diaria (data, capacidade_total_litros, armazenamento_utilizado_litros, usuario_id) VALUES (?, ?, ?, ?)",
                                                 (data, capacidade_total, armazenamento_utilizado, session['usuario_id']), fetch=False)

            database.execute_query("DELETE FROM producao_tanques WHERE producao_diaria_id = ?", (prod_id,), fetch=False)
            for t, td in tanques_data.items():
                database.execute_query("INSERT INTO producao_tanques (producao_diaria_id, tanque, produto, valor_litros) VALUES (?, ?, ?, ?)",
                                       (prod_id, t, td['produto'], td['valor_litros']), fetch=False)

            registrar_log('qualidade', 'Lançou armazenamento diário em formulação')
            mensagem = "Produção do dia salva."
            dia_selecionado = data

        elif 'excluir' in request.form:
            prod_id = request.form.get('id')
            database.execute_query("DELETE FROM producao_diaria WHERE id = ?", (prod_id,), fetch=False)
            registrar_log('qualidade', 'Excluiu um lançamento de formulação')
            mensagem = "Registro excluído."

    producao = database.execute_query_single("SELECT * FROM producao_diaria WHERE data = ?", (dia_selecionado,))
    valores_tanques = {}
    if producao:
        t_rows = database.execute_query("SELECT * FROM producao_tanques WHERE producao_diaria_id = ?", (producao['id'],))
        for tr in t_rows:
            valores_tanques[tr['tanque']] = tr

    historico = database.execute_query("SELECT * FROM producao_diaria ORDER BY data DESC LIMIT 30")

    return render_template('admin_em_formulacao.html', mensagem=mensagem, dia_selecionado=dia_selecionado,
                           tanques_disponiveis=tanques_disponiveis, producao=producao, valores_tanques=valores_tanques,
                           historico=historico)
def registrar_log(modulo, descricao):
    if not is_logged_in(): return
    query = "INSERT INTO log_auditoria (usuario_id, modulo, descricao) VALUES (?, ?, ?)"
    database.execute_query(query, (session['usuario_id'], modulo, descricao), fetch=False)

@app.route('/admin/mural', methods=['GET', 'POST'])
def admin_mural():
    if departamento_usuario() not in ['RH', 'Marketing']:
        return "Acesso negado.", 403

    mensagem = None
    erro = None
    pasta_uploads = os.path.join(os.path.dirname(__file__), 'static', 'uploads', 'mural')

    if request.method == 'POST':
        if 'publicar' in request.form:
            titulo = request.form.get('titulo', '').strip()
            tipo = request.form.get('tipo', 'texto')
            data_inicio = request.form.get('data_inicio')
            data_fim = request.form.get('data_fim')
            conteudo = request.form.get('conteudo', '')
            imagem_nome = None

            if tipo == 'imagem':
                imagem = request.files.get('imagem')
                if imagem and imagem.filename:
                    import uuid
                    ext = imagem.filename.rsplit('.', 1)[1].lower() if '.' in imagem.filename else ''
                    if ext in ['jpg', 'jpeg', 'png']:
                        imagem_nome = f"mural_{uuid.uuid4().hex}.{ext}"
                        imagem.save(os.path.join(pasta_uploads, imagem_nome))
                    else:
                        erro = "Envie apenas arquivos JPG ou PNG."
                if not imagem_nome and not erro:
                    erro = "Selecione uma imagem para publicar."

            if not erro:
                query = """
                    INSERT INTO mural_posts (tipo, titulo, conteudo, imagem_path, departamento_id, data_inicio, data_fim, usuario_id)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """
                database.execute_query(query, (tipo, titulo, conteudo, imagem_nome, session['departamento_id'], data_inicio, data_fim, session['usuario_id']), fetch=False)
                registrar_log('mural', 'Publicou um novo aviso no mural')
                mensagem = "Aviso publicado."

        elif 'alternar_ativo' in request.form:
            post_id = request.form.get('id')
            database.execute_query("UPDATE mural_posts SET ativo = 1 - ativo WHERE id = ?", (post_id,), fetch=False)
            registrar_log('mural', 'Ativou/desativou um aviso')
            mensagem = "Status alterado."

        elif 'excluir' in request.form:
            post_id = request.form.get('id')
            imagem_antiga = request.form.get('imagem_antiga')
            database.execute_query("DELETE FROM mural_posts WHERE id = ?", (post_id,), fetch=False)
            if imagem_antiga:
                p = os.path.join(pasta_uploads, imagem_antiga)
                if os.path.exists(p): os.remove(p)
            registrar_log('mural', 'Excluiu um aviso')
            mensagem = "Aviso excluído."

    mostrar_todos = request.args.get('mostrar') == 'todos'
    hoje = datetime.datetime.now().strftime('%Y-%m-%d')
    if mostrar_todos:
        query = "SELECT * FROM mural_posts ORDER BY criado_em DESC"
    else:
        query = "SELECT * FROM mural_posts WHERE ativo = 1 AND ? BETWEEN data_inicio AND data_fim ORDER BY criado_em DESC"

    mural = database.execute_query(query, () if mostrar_todos else (hoje,))

    return render_template('admin_mural.html', mural=mural, mensagem=mensagem, erro=erro, hoje=hoje, mostrar_todos='1' if mostrar_todos else '0')

@app.route('/admin/noticias', methods=['GET', 'POST'])
def admin_noticias():
    if departamento_usuario() not in ['RH', 'Marketing']:
        return "Acesso negado.", 403

    mensagem = None
    erro = None

    if request.method == 'POST':
        if 'adicionar' in request.form:
            url_feed = request.form.get('url_feed', '').strip()
            if not url_feed.startswith('http'):
                erro = "Informe uma URL de feed válida."
            else:
                database.execute_query("INSERT INTO noticias_rss (url_feed, departamento_id, usuario_id) VALUES (?, ?, ?)",
                                       (url_feed, session['departamento_id'], session['usuario_id']), fetch=False)
                registrar_log('noticias', 'Cadastrou um novo feed de notícias (RSS)')
                mensagem = "Feed cadastrado."

        elif 'editar' in request.form:
            url_feed = request.form.get('url_feed', '').strip()
            feed_id = request.form.get('id')
            if not url_feed.startswith('http'):
                erro = "Informe uma URL de feed válida."
            else:
                database.execute_query("UPDATE noticias_rss SET url_feed = ? WHERE id = ?", (url_feed, feed_id), fetch=False)
                registrar_log('noticias', 'Editou um feed de notícias (RSS)')
                mensagem = "Feed atualizado."

        elif 'alternar_ativo' in request.form:
            feed_id = request.form.get('id')
            database.execute_query("UPDATE noticias_rss SET ativo = 1 - ativo WHERE id = ?", (feed_id,), fetch=False)
            registrar_log('noticias', 'Ativou/desativou um feed de notícias')
            mensagem = "Feed atualizado."

        elif 'excluir' in request.form:
            feed_id = request.form.get('id')
            database.execute_query("DELETE FROM noticias_rss WHERE id = ?", (feed_id,), fetch=False)
            registrar_log('noticias', 'Excluiu um feed de notícias')
            mensagem = "Feed excluído."

    feeds = database.execute_query("SELECT * FROM noticias_rss ORDER BY id DESC")
    return render_template('admin_noticias.html', feeds=feeds, mensagem=mensagem, erro=erro)
@app.route('/admin/usuarios', methods=['GET', 'POST'])
def admin_usuarios():
    if departamento_usuario() != 'TI':
        return "Acesso negado.", 403

    mensagem = None
    erro = None

    def usuario_tem_registros(uid):
        tabelas = {
            "producao_diaria": "usuario_id",
            "status_qualidade_dia": "usuario_id",
            "mural_posts": "usuario_id",
            "indicadores": "usuario_id",
            "indicadores_arquivos": "usuario_id",
            "log_auditoria": "usuario_id"
        }
        for tabela, coluna in tabelas.items():
            check = database.execute_query_single(f"SELECT COUNT(*) as c FROM {tabela} WHERE {coluna} = ?", (uid,))
            if check and check['c'] > 0:
                return True
        return False

    if request.method == 'POST':
        if 'criar' in request.form:
            nome = request.form.get('nome')
            email = request.form.get('email')
            senha = request.form.get('senha')
            dep_id = request.form.get('departamento_id')
            hash_senha = generate_password_hash(senha)

            try:
                database.execute_query("INSERT INTO usuarios (nome, email, senha_hash, departamento_id, deve_trocar_senha) VALUES (?, ?, ?, ?, 1)",
                                       (nome, email, hash_senha, dep_id), fetch=False)
                registrar_log('usuarios', 'Criou um novo usuário')
                mensagem = "Usuário criado. Ele(a) vai precisar trocar a senha no primeiro acesso."
            except Exception as e:
                erro = "Erro ao criar usuário (e-mail já existe?)."

        elif 'redefinir_senha' in request.form:
            uid = request.form.get('id')
            nova_senha = request.form.get('nova_senha')
            hash_senha = generate_password_hash(nova_senha)
            database.execute_query("UPDATE usuarios SET senha_hash = ?, deve_trocar_senha = 1 WHERE id = ?", (hash_senha, uid), fetch=False)
            registrar_log('usuarios', 'Redefiniu a senha de um usuário')
            mensagem = "Senha redefinida."

        elif 'alternar_ativo' in request.form:
            uid = request.form.get('id')
            if uid == str(session['usuario_id']):
                erro = "Você não pode desativar o próprio usuário enquanto está logado."
            else:
                database.execute_query("UPDATE usuarios SET ativo = 1 - ativo WHERE id = ?", (uid,), fetch=False)
                registrar_log('usuarios', 'Ativou/desativou um usuário')
                mensagem = "Status do usuário alterado."

        elif 'excluir' in request.form:
            uid = request.form.get('id')
            if uid == str(session['usuario_id']):
                erro = "Você não pode excluir o próprio usuário."
            elif usuario_tem_registros(uid):
                erro = "Este usuário não pode ser excluído porque já inseriu dados no sistema. Desative-o em vez disso para manter o histórico."
            else:
                database.execute_query("DELETE FROM tokens_lembrar WHERE usuario_id = ?", (uid,), fetch=False)
                database.execute_query("DELETE FROM usuarios WHERE id = ?", (uid,), fetch=False)
                registrar_log('usuarios', 'Excluiu um usuário')
                mensagem = "Usuário excluído definitivamente."

    usuarios = database.execute_query("""
        SELECT u.*, d.nome AS departamento_nome
        FROM usuarios u
        JOIN departamentos d ON d.id = u.departamento_id
        ORDER BY u.nome
    """)
    departamentos = database.execute_query("SELECT * FROM departamentos ORDER BY nome")

    return render_template('admin_usuarios.html', usuarios=usuarios, departamentos=departamentos, mensagem=mensagem, erro=erro)

@app.route('/admin/usuario-editar/<int:id>', methods=['GET', 'POST'])
def admin_usuario_editar(id):
    if departamento_usuario() != 'TI':
        return "Acesso negado.", 403

    mensagem = None
    if request.method == 'POST':
        nome = request.form.get('nome')
        email = request.form.get('email')
        dep_id = request.form.get('departamento_id')
        database.execute_query("UPDATE usuarios SET nome = ?, email = ?, departamento_id = ? WHERE id = ?", (nome, email, dep_id, id), fetch=False)
        registrar_log('usuarios', 'Editou os dados de um usuário')
        mensagem = "Usuário atualizado."

    usuario = database.execute_query_single("SELECT * FROM usuarios WHERE id = ?", (id,))
    if not usuario:
        return redirect(url_for('admin_usuarios'))

    departamentos = database.execute_query("SELECT * FROM departamentos ORDER BY nome")
    return render_template('admin_usuario_editar.html', usuario=usuario, departamentos=departamentos, mensagem=mensagem)

@app.route('/admin/auditoria')
def admin_auditoria():
    if departamento_usuario() != 'TI':
        return "Acesso negado.", 403

    # Limpeza de logs antigos
    database.execute_query("DELETE FROM log_auditoria WHERE data_hora < datetime('now', '-30 days')", fetch=False)

    data_filtro = request.args.get('data', '')

    sql = """
        SELECT la.*, u.nome AS usuario_nome, d.nome AS departamento
        FROM log_auditoria la
        JOIN usuarios u ON u.id = la.usuario_id
        JOIN departamentos d ON d.id = u.departamento_id
    """

    if data_filtro:
        sql += " WHERE la.data_hora LIKE ?"
        sql += " ORDER BY la.data_hora DESC LIMIT 300"
        registros = database.execute_query(sql, (f"{data_filtro}%",))
    else:
        sql += " ORDER BY la.data_hora DESC LIMIT 300"
        registros = database.execute_query(sql)

    return render_template('admin_auditoria.html', registros=registros, data_filtro=data_filtro)
@app.route('/gerar-senha', methods=['GET', 'POST'])
def gerar_senha():
    hash_senha = None
    if request.method == 'POST':
        senha = request.form.get('senha')
        if senha:
            hash_senha = generate_password_hash(senha)
    return render_template('gerar_senha.html', hash=hash_senha)

@app.route('/trocar-senha', methods=['GET', 'POST'])
def trocar_senha():
    if not is_logged_in():
        return redirect(url_for('login'))

    erro = None
    obrigatorio = precisa_trocar_senha()

    if request.method == 'POST':
        nova_senha = request.form.get('nova_senha', '')
        confirmacao = request.form.get('confirmar_senha', '')

        if len(nova_senha) < 6:
            erro = "A nova senha precisa ter pelo menos 6 caracteres."
        elif nova_senha != confirmacao:
            erro = "As senhas não coincidem."
        else:
            hash_senha = generate_password_hash(nova_senha)
            database.execute_query("UPDATE usuarios SET senha_hash = ?, deve_trocar_senha = 0 WHERE id = ?",
                                   (hash_senha, session['usuario_id']), fetch=False)
            session['deve_trocar_senha'] = False
            return redirect(get_destino_apos_login(session.get('departamento_nome')))

    return render_template('trocar_senha.html', erro=erro, obrigatorio=obrigatorio)

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)
